#!/usr/bin/env python3
"""Consolidated Excel export for RK×BG assortment cards."""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

TITLE = "Справочник сценариев и категорий"

NAV_SHEET = "Фильтр"
HELPER_SHEET = "РК и БГ"
LINK_COL = "Перейти к данным"

DATA_SHEETS = frozenset({"ML4×Сценарии", "ML4 свод", "Сценарии свод", "Карточки"})

ROW_TYPE_ML4 = "ML4"
ROW_TYPE_SCEN = "Сценарий"
ROW_TYPE_COMP = "Состав"

HIER_COLUMNS = [
    "Уровень",
    "Lvl1",
    "Lvl2",
    "Lvl3",
    "Lvl4",
    "Мега",
    "Сценарий",
    "Миссия",
    "Домен (MCC)",
    "GMV сценария, ₽",
    "GMV ML4, ₽",
    "GMV сцен. Σ, ₽",
    "Доля ML4 в GMV сценария, %",
    "Эталон 15",
    "Эталон МАКС",
    "Эталон Σ",
    "n сценариев",
    "n_ml4 всего",
    "n_ml4 в БГ",
    "ML4 в этой БГ",
    "Эксклюзив",
    "РК владелец",
    "БГ владелец",
    "Gap Сбер, п.п.",
]

ML4_ROW_FONT = Font(bold=True)
SCEN_ROW_FONT = Font(bold=False)
COMP_ROW_FONT = Font(color="595959")
FOCUS_FILL = PatternFill("solid", fgColor="DDEBF7")
THIS_BG_FILL = PatternFill("solid", fgColor="E2EFDA")
OTHER_BG_FILL = PatternFill("solid", fgColor="F2F2F2")
HIER_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HIER_HEADER_FONT = Font(color="FFFFFF", bold=True)

RK_FILL_COLORS = (
    "E8F4FD",
    "FFF2CC",
    "E2EFDA",
    "FCE4D6",
    "EDEDED",
    "D9E1F2",
    "F4CCCC",
    "D9D2E9",
    "CFE2F3",
)

HYPERLINK_FONT = Font(color="0563C1", underline="single")
NAV_BACK_LABEL = "← Содержание"
NAV_BACK_FONT = Font(color="0563C1", underline="single", bold=True)
NAV_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
NAV_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _flag_etalon_gt0(val: object) -> bool:
    try:
        return int(float(val or 0)) > 0
    except (TypeError, ValueError):
        return False


def _norm_ml4(name: object) -> str:
    return str(name or "").strip().lower().replace("ё", "е")


def _etalon_full_sum(m: dict) -> int:
    v = m.get("etalon_full_sum")
    if v is not None:
        try:
            return int(float(v))
        except (TypeError, ValueError):
            pass
    return int(m.get("etalon_sum") or 0)


def _has_etalon_sku(m: dict) -> bool:
    return _etalon_full_sum(m) > 0


def _allocate_gmv(gmv: float, weights: list[float]) -> list[float]:
    if not weights:
        return []
    positive = [max(float(w or 0.0), 0.0) for w in weights]
    total = float(sum(positive))
    if total <= 0:
        share = float(gmv) / len(weights)
        return [share] * len(weights)
    return [float(gmv) * (w / total) for w in positive]


def _scenario_by_mission(card: dict, mission: str) -> dict | None:
    for s in card.get("scenarios") or []:
        if str(s.get("mission") or "") == mission:
            return s
    return None


def _scenario_ml4_rows(scenario: dict) -> list[dict]:
    all_rows = scenario.get("ml4_all")
    if all_rows:
        return list(all_rows)
    return list(scenario.get("ml4") or [])


def _ml4_gmv_split(scenario: dict, focus_lvl4: str | None) -> list[dict]:
    rows = _scenario_ml4_rows(scenario)
    scen_gmv = float(scenario.get("gmv") or 0.0)
    has_ml4 = any(m.get("gmv_ml4") is not None for m in rows)
    if has_ml4:
        allocs = [float(m.get("gmv_ml4") or 0.0) for m in rows]
    else:
        allocs = _allocate_gmv(scen_gmv, [_etalon_full_sum(m) for m in rows])
    total_alloc = float(sum(allocs))
    focus = _norm_ml4(focus_lvl4) if focus_lvl4 else ""
    out: list[dict] = []
    for m, alloc in zip(rows, allocs):
        rec = dict(m)
        rec["gmv_allocated"] = alloc
        rec["gmv_share"] = (alloc / total_alloc) if total_alloc > 0 else 0.0
        rec["is_focus"] = bool(focus) and _norm_ml4(m.get("lvl4")) == focus
        out.append(rec)
    return out


def _empty_hier_row() -> dict:
    return {col: None for col in HIER_COLUMNS}


def _hier_row(**kwargs) -> dict:
    row = _empty_hier_row()
    row.update(kwargs)
    return row


def build_hierarchical_card_sheet(
    card: dict,
    *,
    mcc_gap: dict[str, float] | None = None,
) -> tuple[pd.DataFrame, list[tuple[int, int, int]], list[dict]]:
    """ML4 → сценарии → состав сценария; returns (df, excel row group ranges, row styles)."""
    gap_map = {str(k).strip(): float(v) for k, v in (mcc_gap or {}).items() if str(k).strip()}
    rows: list[dict] = []
    row_styles: list[dict] = []
    group_ranges: list[tuple[int, int, int]] = []
    excel_row = 2  # row 1 = header

    ml4_rows = list(card.get("ml4s") or [])
    ml4_rows.sort(
        key=lambda m: (
            -float(m.get("gmv_proxy") or 0.0),
            -int(m.get("etalon_sum") or 0),
            str(m.get("lvl4") or ""),
        )
    )

    for m in ml4_rows:
        if not _has_etalon_sku(m) and not m.get("missions"):
            continue
        owner = m.get("exclusive_owner") or ""
        if not owner and m.get("in_unused"):
            owner = "вне сценариев"
        rows.append(
            _hier_row(
                **{
                    "Уровень": ROW_TYPE_ML4,
                    "Lvl1": m.get("lvl1") or "",
                    "Lvl2": m.get("lvl2") or "",
                    "Lvl3": m.get("lvl3") or "",
                    "Lvl4": m.get("lvl4") or "",
                    "GMV ML4, ₽": float(m.get("gmv_proxy") or 0.0) or None,
                    "GMV сцен. Σ, ₽": float(m.get("gmv_touch") or 0.0) or None,
                    "Эталон 15": int(m.get("etalon_15") or 0),
                    "Эталон МАКС": int(m.get("etalon_max") or 0),
                    "Эталон Σ": int(m.get("etalon_sum") or 0),
                    "n сценариев": int(m.get("n_scenarios") or 0),
                    "Эксклюзив": owner,
                }
            )
        )
        row_styles.append({"level": ROW_TYPE_ML4})
        excel_row += 1

        missions = list(m.get("missions") or [])
        scen_entries: list[dict] = []
        for mission in missions:
            scen = _scenario_by_mission(card, mission)
            scen_entries.append(scen if scen else {"mission": mission})
        scen_entries.sort(
            key=lambda s: (
                -float(s.get("gmv") or 0.0),
                -int(s.get("etalon_sum") or 0),
                str(s.get("mission") or ""),
            )
        )

        if not scen_entries:
            continue

        block_start = excel_row
        focus_lvl4 = str(m.get("lvl4") or "")

        for scen in scen_entries:
            mission = str(scen.get("mission") or "")
            mcc = str(scen.get("mcc_group") or "")
            gap_pp = gap_map.get(mcc)
            scen_gmv = float(scen.get("gmv") or 0.0)
            split_all = _ml4_gmv_split(scen, focus_lvl4) if scen.get("ml4_all") or scen.get("ml4") else []
            split = [x for x in split_all if _has_etalon_sku(x)]
            focus = next((x for x in split_all if x.get("is_focus")), None)
            if focus is None:
                focus = next((x for x in split if x.get("is_focus")), None)
            focus_gmv = float(focus.get("gmv_allocated") or 0.0) if focus else 0.0
            focus_share = float(focus.get("gmv_share") or 0.0) if focus else 0.0

            rows.append(
                _hier_row(
                    **{
                        "Уровень": ROW_TYPE_SCEN,
                        "Мега": scen.get("mega") or "",
                        "Сценарий": scen.get("scenario") or mission,
                        "Миссия": mission,
                        "Домен (MCC)": mcc,
                        "GMV сценария, ₽": scen_gmv if scen_gmv else None,
                        "GMV ML4, ₽": focus_gmv if focus_gmv else None,
                        "Доля ML4 в GMV сценария, %": round(focus_share * 100, 4) if focus else None,
                        "Эталон 15": int(focus.get("etalon_15") or 0) if focus else None,
                        "Эталон МАКС": int(focus.get("etalon_max") or 0) if focus else None,
                        "Эталон Σ": int(focus.get("etalon_sum") or 0) if focus else None,
                        "n_ml4 всего": int(scen.get("n_ml4_all") or len(split_all) or 0),
                        "n_ml4 в БГ": int(scen.get("n_ml4") or 0),
                        "Gap Сбер, п.п.": round(gap_pp * 100, 2) if gap_pp is not None else None,
                    }
                )
            )
            row_styles.append({"level": ROW_TYPE_SCEN})
            excel_row += 1

            comp_start = excel_row
            for x in split:
                gmv_a = float(x.get("gmv_allocated") or 0.0)
                share = float(x.get("gmv_share") or 0.0)
                in_bg = bool(x.get("in_this_bg"))
                is_focus = bool(x.get("is_focus"))
                marker = ""
                if is_focus:
                    marker = "эта"
                elif in_bg:
                    marker = "да"
                rows.append(
                    _hier_row(
                        **{
                            "Уровень": ROW_TYPE_COMP,
                            "Lvl1": x.get("lvl1") or "",
                            "Lvl2": x.get("lvl2") or "",
                            "Lvl3": x.get("lvl3") or "",
                            "Lvl4": x.get("lvl4") or "",
                            "GMV ML4, ₽": gmv_a if gmv_a else None,
                            "Доля ML4 в GMV сценария, %": round(share * 100, 4) if share else None,
                            "Эталон 15": int(x.get("etalon_15") or 0),
                            "Эталон МАКС": int(x.get("etalon_max") or 0),
                            "Эталон Σ": int(x.get("etalon_sum") or 0),
                            "ML4 в этой БГ": marker,
                            "Эксклюзив": "да" if x.get("exclusive") else "",
                            "РК владелец": x.get("rk") or "",
                            "БГ владелец": x.get("bg") or "",
                        }
                    )
                )
                row_styles.append(
                    {
                        "level": ROW_TYPE_COMP,
                        "is_focus": is_focus,
                        "in_this_bg": in_bg and not is_focus,
                    }
                )
                excel_row += 1

            comp_end = excel_row - 1
            if comp_end >= comp_start:
                group_ranges.append((comp_start, comp_end, 2))

        block_end = excel_row - 1
        if block_end >= block_start:
            group_ranges.append((block_start, block_end, 1))

    df = pd.DataFrame(rows, columns=HIER_COLUMNS) if rows else pd.DataFrame(columns=HIER_COLUMNS)
    return df, group_ranges, row_styles


def _sanitize_sheet_name(name: str) -> str:
    for ch in r'\/:*?[]':
        name = name.replace(ch, "_")
    return name.strip()[:31]


def _rk_short(rk: str) -> str:
    rk = str(rk).strip()
    if not rk:
        return rk
    parts = rk.split()
    if len(parts) >= 2 and len(rk) > 20:
        return parts[0]
    return rk


def _card_sheet_name(rk: str, bg: str, used: set[str]) -> str:
    base = _sanitize_sheet_name(f"{_rk_short(rk)} · {bg}")
    name = base or "Карточка"
    idx = 2
    while name in used:
        suffix = f" ({idx})"
        name = _sanitize_sheet_name(base[: max(1, 31 - len(suffix))] + suffix)
        idx += 1
    used.add(name)
    return name


def _card_sheet_map(cards_df: pd.DataFrame) -> dict[tuple[str, str], str]:
    used: set[str] = set(DATA_SHEETS) | {NAV_SHEET, HELPER_SHEET, "Справка"}
    mapping: dict[tuple[str, str], str] = {}
    if cards_df.empty:
        return mapping
    for _, row in cards_df.iterrows():
        rk = str(row.get("РК") or "")
        bg = str(row.get("БГ") or "")
        mapping[(rk, bg)] = _card_sheet_name(rk, bg, used)
    return mapping


def build_ml4_scenario_sheet(
    assortment: dict,
    *,
    mcc_gap: dict[str, float] | None = None,
) -> pd.DataFrame:
    """One row per (RK, BG, mission, ML4) — full scenario composition with BG flag."""
    gap_map = {str(k).strip(): float(v) for k, v in (mcc_gap or {}).items() if str(k).strip()}
    rows: list[dict] = []
    cards = assortment.get("cards") or {}
    for card in cards.values():
        rk = str(card.get("rk") or "")
        bg = str(card.get("bg") or "")
        for scen in card.get("scenarios") or []:
            mission = str(scen.get("mission") or "")
            mega = str(scen.get("mega") or "")
            scenario = str(scen.get("scenario") or "")
            mcc = str(scen.get("mcc_group") or "")
            scen_gmv = float(scen.get("gmv") or 0.0)
            leaf_share = float(scen.get("leaf_share") or 0.0)
            gap_pp = gap_map.get(mcc)
            scen_e15 = int(scen.get("etalon_15") or 0)
            scen_emax = int(scen.get("etalon_max") or 0)
            scen_esum = int(scen.get("etalon_sum") or 0)
            for m in scen.get("ml4_all") or scen.get("ml4") or []:
                gmv_ml4 = m.get("gmv_ml4")
                try:
                    gmv_ml4_f = float(gmv_ml4) if gmv_ml4 is not None else None
                except (TypeError, ValueError):
                    gmv_ml4_f = None
                share = (gmv_ml4_f / scen_gmv) if scen_gmv > 0 and gmv_ml4_f is not None else None
                et_sum = int(m.get("etalon_sum") or 0)
                rows.append(
                    {
                        "РК": rk,
                        "БГ": bg,
                        "Домен (MCC)": mcc,
                        "Lvl1": m.get("lvl1") or "",
                        "Lvl2": m.get("lvl2") or "",
                        "Lvl3": m.get("lvl3") or "",
                        "Lvl4": m.get("lvl4") or "",
                        "Мега": mega,
                        "Сценарий": scenario,
                        "Миссия": mission,
                        "GMV сценария, ₽": scen_gmv if scen_gmv else None,
                        "Доля сценария в кошельке, %": round(leaf_share * 100, 4) if leaf_share else None,
                        "GMV ML4 в сценарии, ₽": gmv_ml4_f,
                        "Доля ML4 в GMV сценария, %": round(share * 100, 4) if share is not None else None,
                        "Эталон 15 (доля)": int(m.get("etalon_15") or 0),
                        "Эталон МАКС (доля)": int(m.get("etalon_max") or 0),
                        "Эталон Σ (доля)": et_sum,
                        "Эталон 15 (полный)": int(m.get("etalon_full_15") or m.get("etalon_15") or 0),
                        "Эталон МАКС (полный)": int(m.get("etalon_full_max") or m.get("etalon_max") or 0),
                        "Эталон Σ (полный)": int(m.get("etalon_full_sum") or m.get("etalon_sum") or 0),
                        "Эталон сценария 15 (экскл. БГ)": scen_e15,
                        "Эталон сценария МАКС (экскл. БГ)": scen_emax,
                        "Эталон сценария Σ (экскл. БГ)": scen_esum,
                        "Эксклюзив сценария": "да" if m.get("exclusive") else "",
                        "Владелец (если не экскл.)": m.get("owner") or "",
                        "ML4 в этой БГ": "да" if m.get("in_this_bg") else "",
                        "Эталон Σ > 0": "да" if _flag_etalon_gt0(et_sum) else "",
                        "Gap Сбер, п.п.": round(gap_pp * 100, 2) if gap_pp is not None else None,
                        "n_ml4 в БГ": int(scen.get("n_ml4") or 0),
                        "n_ml4 всего": int(scen.get("n_ml4_all") or scen.get("n_ml4") or 0),
                    }
                )
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    return df.sort_values(
        ["РК", "БГ", "GMV сценария, ₽", "Эталон Σ (полный)", "Lvl4", "Миссия"],
        ascending=[True, True, False, False, True, True],
        na_position="last",
    ).reset_index(drop=True)


def build_ml4_summary_sheet(assortment: dict) -> pd.DataFrame:
    """Aggregated ML4 index per RK×BG card."""
    rows: list[dict] = []
    for card in (assortment.get("cards") or {}).values():
        rk = str(card.get("rk") or "")
        bg = str(card.get("bg") or "")
        for m in card.get("ml4s") or []:
            et_sum = int(m.get("etalon_sum") or 0)
            missions = m.get("missions") or []
            rows.append(
                {
                    "РК": rk,
                    "БГ": bg,
                    "Lvl1": m.get("lvl1") or "",
                    "Lvl2": m.get("lvl2") or "",
                    "Lvl3": m.get("lvl3") or "",
                    "Lvl4": m.get("lvl4") or "",
                    "БГ категории": m.get("bg") or bg,
                    "Эталон 15": int(m.get("etalon_15") or 0),
                    "Эталон МАКС": int(m.get("etalon_max") or 0),
                    "Эталон Σ": et_sum,
                    "Эталон Σ > 0": "да" if _flag_etalon_gt0(et_sum) else "",
                    "n сценариев": int(m.get("n_scenarios") or 0),
                    "GMV touch (сумма GMV сценариев), ₽": float(m.get("gmv_touch") or 0) or None,
                    "GMV ML4 (сумма по сценариям), ₽": float(m.get("gmv_proxy") or 0) or None,
                    "Экскл. сценарий": m.get("exclusive_owner") or "",
                    "Вне сценариев": "да" if m.get("in_unused") else "",
                    "Сценарии": " · ".join(missions[:8]) + ("…" if len(missions) > 8 else ""),
                }
            )
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    return df.sort_values(
        ["РК", "БГ", "GMV ML4 (сумма по сценариям), ₽", "Эталон Σ", "Lvl4"],
        ascending=[True, True, False, False, True],
        na_position="last",
    ).reset_index(drop=True)


def build_scenario_summary_sheet(
    assortment: dict,
    *,
    mcc_gap: dict[str, float] | None = None,
) -> pd.DataFrame:
    """Scenario-level rollup per RK×BG card."""
    gap_map = {str(k).strip(): float(v) for k, v in (mcc_gap or {}).items() if str(k).strip()}
    rows: list[dict] = []
    for card in (assortment.get("cards") or {}).values():
        rk = str(card.get("rk") or "")
        bg = str(card.get("bg") or "")
        for s in card.get("scenarios") or []:
            mcc = str(s.get("mcc_group") or "")
            gap_pp = gap_map.get(mcc)
            esum = int(s.get("etalon_sum") or 0)
            rows.append(
                {
                    "РК": rk,
                    "БГ": bg,
                    "Мега": s.get("mega") or "",
                    "Сценарий": s.get("scenario") or "",
                    "Миссия": s.get("mission") or "",
                    "Домен (MCC)": mcc,
                    "GMV, ₽": float(s.get("gmv") or 0) or None,
                    "Доля в кошельке, %": round(float(s.get("leaf_share") or 0) * 100, 4) or None,
                    "n_ml4 в БГ": int(s.get("n_ml4") or 0),
                    "n_ml4 всего": int(s.get("n_ml4_all") or s.get("n_ml4") or 0),
                    "Эталон 15 (экскл. БГ)": int(s.get("etalon_15") or 0),
                    "Эталон МАКС (экскл. БГ)": int(s.get("etalon_max") or 0),
                    "Эталон Σ (экскл. БГ)": esum,
                    "Эталон Σ > 0": "да" if _flag_etalon_gt0(esum) else "",
                    "Статус эталона": s.get("etalon_status") or "",
                    "Потеряно SKU (не экскл.)": int(s.get("n_lost") or 0),
                    "Gap Сбер, п.п.": round(gap_pp * 100, 2) if gap_pp is not None else None,
                }
            )
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    return df.sort_values(
        ["РК", "БГ", "GMV, ₽", "Миссия"],
        ascending=[True, True, False, True],
        na_position="last",
    ).reset_index(drop=True)


def build_cards_sheet(assortment: dict) -> pd.DataFrame:
    """RK×BG card headers."""
    rows: list[dict] = []
    for card in (assortment.get("cards") or {}).values():
        esum = int(card.get("etalon_sum") or 0)
        rows.append(
            {
                "РК": card.get("rk") or "",
                "БГ": card.get("bg") or "",
                "n сценариев": int(card.get("n_scenarios") or 0),
                "n ML4": int(card.get("n_ml4") or 0),
                "n ML4 в сценариях": int(card.get("n_ml4_in_scenarios") or 0),
                "Эталон 15": int(card.get("etalon_15") or 0),
                "Эталон МАКС": int(card.get("etalon_max") or 0),
                "Эталон Σ": esum,
                "Эталон Σ > 0": "да" if _flag_etalon_gt0(esum) else "",
                "Вне сценариев, n": int(card.get("unused_n") or 0),
                "Вне сценариев, Σ": int(card.get("unused_sum") or 0),
            }
        )
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    return df.sort_values(["РК", "БГ"]).reset_index(drop=True)


def build_navigation_sheet(
    cards_df: pd.DataFrame,
    card_sheet_map: dict[tuple[str, str], str],
) -> pd.DataFrame:
    rows: list[dict] = []
    for _, row in cards_df.iterrows():
        rk = str(row.get("РК") or "")
        bg = str(row.get("БГ") or "")
        sheet = card_sheet_map.get((rk, bg), "")
        rows.append(
            {
                "РК": rk,
                "БГ": bg,
                "n сценариев": row.get("n сценариев"),
                "n ML4": row.get("n ML4"),
                "Эталон Σ": row.get("Эталон Σ"),
                LINK_COL: f"Открыть → {sheet}" if sheet else "",
                "_sheet": sheet,
            }
        )
    return pd.DataFrame(rows)


def build_rk_bg_helper_sheet(ml4_df: pd.DataFrame) -> pd.DataFrame:
    rks = sorted(str(x) for x in ml4_df["РК"].dropna().unique()) if not ml4_df.empty else []
    bgs = sorted(str(x) for x in ml4_df["БГ"].dropna().unique()) if not ml4_df.empty else []
    max_len = max(len(rks), len(bgs), 1)
    rows: list[dict] = []
    for i in range(max_len):
        rows.append(
            {
                "РК (уникальные)": rks[i] if i < len(rks) else "",
                "БГ (уникальные)": bgs[i] if i < len(bgs) else "",
            }
        )
    return pd.DataFrame(rows)


def build_spravka_sheet(
    assortment: dict,
    *,
    city: str = "",
    variant: str = "",
    card_sheet_map: dict[tuple[str, str], str] | None = None,
) -> pd.DataFrame:
    meta = assortment.get("snapshot_meta") or {}
    n_cards = len(card_sheet_map or {})
    rows = [
        {"раздел": "название", "ключ": "title", "значение": TITLE, "комментарий": ""},
        {"раздел": "город", "ключ": "city", "значение": city, "комментарий": "только Москва"},
        {"раздел": "вариант", "ключ": "variant", "значение": variant, "комментарий": ""},
        {"раздел": "метод", "ключ": "note", "значение": str(assortment.get("note") or ""), "комментарий": ""},
    ]
    for k, v in sorted(meta.items()):
        rows.append({"раздел": "meta", "ключ": k, "значение": str(v), "комментарий": ""})
    rows.extend(
        [
            {
                "раздел": "фильтр РК×БГ",
                "ключ": "способ 1 — навигация",
                "значение": f"Лист «{NAV_SHEET}»: кликните «{LINK_COL}» в строке своей карточки",
                "комментарий": f"отдельный лист на каждую из {n_cards} карточек",
            },
            {
                "раздел": "фильтр РК×БГ",
                "ключ": "способ 2 — лист карточки",
                "значение": "Листы «Фамилия · БГ»: иерархия ML4 → сценарии → состав",
                "комментарий": "группы строк сворачиваются как в HTML-отчёте",
            },
            {
                "раздел": "фильтр РК×БГ",
                "ключ": "назад к содержанию",
                "значение": f"На листе карточки: «{NAV_BACK_LABEL}» в строке 1 — возврат на «{NAV_SHEET}»",
                "комментарий": "ссылка закреплена над шапкой таблицы",
            },
            {
                "раздел": "фильтр РК×БГ",
                "ключ": "способ 3 — автофильтр",
                "значение": "Лист «ML4×Сценарии»: плоская таблица для сводного просмотра",
                "комментарий": "фильтр по «РК» и «БГ» в шапке таблицы",
            },
            {
                "раздел": "иерархия",
                "ключ": "уровень ML4",
                "значение": "Строка «ML4»: категория бизнес-группы с полным эталоном и GMV",
                "комментарий": "жирный шрифт; клик «−» слева — свернуть сценарии",
            },
            {
                "раздел": "иерархия",
                "ключ": "уровень Сценарий",
                "значение": "Строка «Сценарий»: мега → сценарий, GMV сценария и GMV этой ML4",
                "комментарий": "вложенная группа; «−» — свернуть состав ML4",
            },
            {
                "раздел": "иерархия",
                "ключ": "уровень Состав",
                "значение": "Строка «Состав»: все ML4 сценария (ml4_all); «эта» — родительская ML4",
                "комментарий": "«да» в «ML4 в этой БГ» — другие категории текущей БГ",
            },
            {
                "раздел": "иерархия",
                "ключ": "свернуть/развернуть",
                "значение": "Данные → Группировка → Скрыть детали (ур. 1 или 2) / Показать детали",
                "комментарий": "или кнопки «−» / «+» слева от номеров строк",
            },
            {
                "раздел": "фильтр РК×БГ",
                "ключ": "справочник значений",
                "значение": f"Лист «{HELPER_SHEET}» — списки уникальных РК и БГ",
                "комментарий": "",
            },
            {
                "раздел": "фильтр РК×БГ",
                "ключ": "срезы",
                "значение": "«ML4 свод» и «Сценарии свод» — автофильтр по «РК» и «БГ»",
                "комментарий": "",
            },
            {
                "раздел": "листы",
                "ключ": NAV_SHEET,
                "значение": "навигация по карточкам РК×БГ с гиперссылками",
                "комментарий": "начните с этого листа",
            },
            {
                "раздел": "листы",
                "ключ": "карточки РК×БГ",
                "значение": "основной вид: ML4 → сценарии → полный состав сценария",
                "комментарий": f"один лист на пару РК×БГ; «{NAV_BACK_LABEL}» в A1 — к содержанию",
            },
            {
                "раздел": "листы",
                "ключ": "ML4×Сценарии",
                "значение": "плоская таблица для power users и сводного фильтра",
                "комментарий": "колонка «ML4 в этой БГ»; строки подсвечены по РК",
            },
            {
                "раздел": "листы",
                "ключ": "ML4 свод",
                "значение": "индекс категорий карточки с полным эталоном",
                "комментарий": "",
            },
            {
                "раздел": "листы",
                "ключ": "Сценарии свод",
                "значение": "сценарии карточки с эксклюзивным эталоном БГ",
                "комментарий": "",
            },
            {
                "раздел": "листы",
                "ключ": "Карточки",
                "значение": "сводные метрики по каждой паре РК×БГ",
                "комментарий": "",
            },
        ]
    )
    return pd.DataFrame(rows)


def _header_map(ws, *, header_row: int = 1) -> dict[str, int]:
    return {str(c.value): c.column for c in ws[header_row] if c.value is not None}


def _apply_excel_table(ws, *, freeze: str = "A2", table_style: str = "TableStyleMedium2", table_id: int = 0) -> None:
    if ws.max_row < 2 or ws.max_column < 1:
        return
    ref = ws.dimensions
    slug = re.sub(r"[^0-9A-Za-z_]", "_", ws.title)[:12] or "Sheet"
    tbl_name = f"Tbl_{table_id}_{slug}"[:30]
    if not tbl_name[0].isalpha():
        tbl_name = "T_" + tbl_name[1:]
    ws.tables.clear()
    tbl = Table(displayName=tbl_name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name=table_style,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tbl)
    ws.freeze_panes = freeze


def _add_back_to_contents_nav(ws) -> tuple[int, int]:
    """Insert nav row at top; return (header_row, data_start_row)."""
    ws.insert_rows(1)
    nav_cell = ws.cell(row=1, column=1)
    nav_cell.value = NAV_BACK_LABEL
    nav_cell.hyperlink = f"#'{NAV_SHEET}'!A1"
    nav_cell.font = NAV_BACK_FONT
    nav_cell.alignment = Alignment(vertical="center")
    merge_end = min(5, ws.max_column)
    if merge_end > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=merge_end)
    ws.row_dimensions[1].height = 20
    return 2, 3


def _style_hierarchical_sheet(
    ws,
    *,
    group_ranges: list[tuple[int, int, int]],
    row_styles: list[dict],
) -> None:
    """Apply outline grouping, fonts, indents, and row fills for card hierarchy."""
    if ws.max_row < 1:
        return

    header_row, data_start_row = _add_back_to_contents_nav(ws)
    group_ranges = [(start + 1, end + 1, level) for start, end, level in group_ranges]

    for cell in ws[header_row]:
        cell.fill = HIER_HEADER_FILL
        cell.font = HIER_HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = f"A{data_start_row}"

    if ws.max_row < data_start_row:
        return

    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.applyStyles = True

    for start, end, level in sorted(group_ranges, key=lambda x: x[2]):
        if end >= start:
            hidden = level >= 2
            ws.row_dimensions.group(start, end, outline_level=level, hidden=hidden)

    headers = _header_map(ws, header_row=header_row)
    lvl_col = headers.get("Уровень", 1)
    lvl4_col = headers.get("Lvl4")
    bg_col = headers.get("ML4 в этой БГ")
    style_cols = sorted({lvl_col, lvl4_col or 0, bg_col or 0} - {0})

    for row_idx, style in enumerate(row_styles, start=data_start_row):
        level = style.get("level")
        indent = 0
        font = SCEN_ROW_FONT
        fill = None
        if level == ROW_TYPE_ML4:
            font = ML4_ROW_FONT
        elif level == ROW_TYPE_SCEN:
            indent = 1
        elif level == ROW_TYPE_COMP:
            indent = 2
            font = COMP_ROW_FONT
            if style.get("is_focus"):
                fill = FOCUS_FILL
            elif style.get("in_this_bg"):
                fill = THIS_BG_FILL
            else:
                fill = OTHER_BG_FILL

        lvl_cell = ws.cell(row=row_idx, column=lvl_col)
        lvl_cell.font = font
        lvl_cell.alignment = Alignment(indent=indent, vertical="top", wrap_text=True)
        if fill is not None and lvl4_col:
            for col_idx in style_cols:
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = font
                cell.fill = fill
        elif level == ROW_TYPE_ML4 and lvl4_col:
            ws.cell(row=row_idx, column=lvl4_col).font = font

    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    )


def _column_widths_from_df(df: pd.DataFrame, *, limit: int = 200) -> list[float]:
    widths: list[float] = []
    sample = df.head(limit)
    for col in df.columns:
        vals = sample[col].fillna("").astype(str)
        maxlen = max(len(str(col)), int(vals.str.len().max()) if len(vals) else 0)
        widths.append(min(52, max(12, maxlen + 2)))
    return widths


def _style_workbook(
    path: Path,
    *,
    card_sheet_names: set[str],
    nav_link_targets: dict[int, str],
    card_sheet_meta: dict[str, dict] | None = None,
    sheet_widths: dict[str, list[float]] | None = None,
) -> None:
    """Autofit columns, Excel tables, hyperlinks, RK row colors."""
    card_sheet_meta = card_sheet_meta or {}
    sheet_widths = sheet_widths or {}
    wb = load_workbook(path)

    for table_id, ws in enumerate(wb.worksheets, start=1):
        preset = sheet_widths.get(ws.title)
        if preset:
            for i, width in enumerate(preset, 1):
                ws.column_dimensions[get_column_letter(i)].width = width
        else:
            for i, col in enumerate(ws.columns, 1):
                maxlen = 0
                for cell in col[:250]:
                    maxlen = max(maxlen, len(str(cell.value)) if cell.value is not None else 0)
                ws.column_dimensions[get_column_letter(i)].width = min(52, max(12, maxlen + 2))

        if ws.title == NAV_SHEET:
            _apply_excel_table(ws, freeze="A2", table_style="TableStyleMedium9", table_id=table_id)
            headers = _header_map(ws)
            link_col = headers.get(LINK_COL)
            if link_col:
                for row_idx, target in nav_link_targets.items():
                    cell = ws.cell(row=row_idx, column=link_col)
                    cell.value = f"Открыть → {target}"
                    cell.hyperlink = f"#'{target}'!A1"
                    cell.font = HYPERLINK_FONT
            for cell in ws[1]:
                cell.fill = NAV_HEADER_FILL
                cell.font = NAV_HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            continue

        if ws.title in card_sheet_names:
            meta = card_sheet_meta.get(ws.title) or {}
            _style_hierarchical_sheet(
                ws,
                group_ranges=meta.get("group_ranges") or [],
                row_styles=meta.get("row_styles") or [],
            )
            continue

        if ws.title in DATA_SHEETS:
            _apply_excel_table(ws, freeze="C2", table_id=table_id)
            if ws.title == "ML4×Сценарии":
                headers = _header_map(ws)
                rk_col = headers.get("РК")
                bg_col = headers.get("БГ")
                if rk_col and ws.max_row > 1:
                    rk_colors: dict[str, str] = {}
                    for row_idx in range(2, ws.max_row + 1):
                        rk_val = str(ws.cell(row=row_idx, column=rk_col).value or "")
                        if rk_val not in rk_colors:
                            rk_colors[rk_val] = RK_FILL_COLORS[len(rk_colors) % len(RK_FILL_COLORS)]
                        fill = PatternFill("solid", fgColor=rk_colors[rk_val])
                        ws.cell(row=row_idx, column=rk_col).fill = fill
                        if bg_col:
                            ws.cell(row=row_idx, column=bg_col).fill = fill
            continue

        if ws.max_row >= 2 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = "A2"

    wb.save(path)


def export_domain_wallet_excel(
    assortment: dict,
    out_path: Path,
    *,
    city: str = "Москва",
    variant: str = "default",
    mcc_gap: dict[str, float] | None = None,
) -> Path:
    """Write consolidated RK card workbook."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    ml4_scenarios = build_ml4_scenario_sheet(assortment, mcc_gap=mcc_gap)
    cards = build_cards_sheet(assortment)
    card_sheet_map = _card_sheet_map(cards)
    cards_raw = assortment.get("cards") or {}

    card_sheets: dict[str, pd.DataFrame] = {}
    card_sheet_meta: dict[str, dict] = {}
    sheet_widths: dict[str, list[float]] = {}
    nav_link_targets: dict[int, str] = {}
    if card_sheet_map:
        for (rk, bg), sheet_name in card_sheet_map.items():
            card_key = f"{rk}||{bg}"
            card = cards_raw.get(card_key)
            if not card:
                for ck, c in cards_raw.items():
                    if str(c.get("rk") or "") == rk and str(c.get("bg") or "") == bg:
                        card = c
                        break
            if not card:
                continue
            hier_df, group_ranges, row_styles = build_hierarchical_card_sheet(card, mcc_gap=mcc_gap)
            card_sheets[sheet_name] = hier_df
            card_sheet_meta[sheet_name] = {
                "group_ranges": group_ranges,
                "row_styles": row_styles,
            }
            sheet_widths[sheet_name] = _column_widths_from_df(hier_df)

    nav_df = build_navigation_sheet(cards, card_sheet_map)
    if not nav_df.empty and "_sheet" in nav_df.columns:
        for i, sheet in enumerate(nav_df["_sheet"].tolist(), start=2):
            if sheet:
                nav_link_targets[i] = str(sheet)
        nav_df = nav_df.drop(columns=["_sheet"])

    sheets: list[tuple[str, pd.DataFrame]] = [(NAV_SHEET, nav_df)]
    sheets.extend((name, df) for name, df in sorted(card_sheets.items(), key=lambda x: x[0]))
    sheets.extend(
        [
            ("ML4×Сценарии", ml4_scenarios),
            ("ML4 свод", build_ml4_summary_sheet(assortment)),
            ("Сценарии свод", build_scenario_summary_sheet(assortment, mcc_gap=mcc_gap)),
            ("Карточки", cards),
            (HELPER_SHEET, build_rk_bg_helper_sheet(ml4_scenarios)),
            (
                "Справка",
                build_spravka_sheet(
                    assortment,
                    city=city,
                    variant=variant,
                    card_sheet_map=card_sheet_map,
                ),
            ),
        ]
    )

    with pd.ExcelWriter(out_path, engine="openpyxl") as xl:
        for name, df in sheets:
            df.to_excel(xl, sheet_name=name[:31], index=False)

    _style_workbook(
        out_path,
        card_sheet_names=set(card_sheets),
        nav_link_targets=nav_link_targets,
        card_sheet_meta=card_sheet_meta,
        sheet_widths=sheet_widths,
    )
    return out_path


if __name__ == "__main__":
    import argparse
    import json
    import sys

    ROOT = Path(__file__).resolve().parent.parent
    ap = argparse.ArgumentParser(description="Export domain_wallet consolidated xlsx from payload json")
    ap.add_argument(
        "--payload",
        default=str(ROOT / "outputs" / "domain_wallet_payload.json"),
        help="domain_wallet_payload.json path",
    )
    ap.add_argument(
        "--out",
        default=str(ROOT / "outputs" / "domain_wallet_consolidated.xlsx"),
    )
    args = ap.parse_args()
    payload = json.loads(Path(args.payload).read_text(encoding="utf-8"))
    assortment = payload.get("assortment") or {}
    gap_map = {
        str(r.get("mcc_group") or "").strip(): float(r.get("gap_pp") or 0.0)
        for r in payload.get("master") or []
        if str(r.get("mcc_group") or "").strip()
    }
    out = export_domain_wallet_excel(
        assortment,
        Path(args.out),
        city=str(payload.get("city") or "Москва"),
        variant=str(payload.get("variant") or ""),
        mcc_gap=gap_map,
    )
    print(f"Wrote {out}")
    sys.exit(0)
